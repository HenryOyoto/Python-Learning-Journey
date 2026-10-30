
from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

# Accept any text from the user
text = input("Enter the heading text: ")

# Create a Word document
doc = Document()

# Configure Heading 1 style
heading = doc.styles["Heading 1"]
heading.font.name = "Times New Roman"
heading.font.size = Pt(14)
heading.font.bold = True
heading.font.color.rgb = RGBColor(0, 31, 63)
heading._element.rPr.rFonts.set(
    qn("w:eastAsia"), "Times New Roman"
)
heading.paragraph_format.line_spacing = 1.5

# Create a multilevel numbering definition
numbering = doc.part.numbering_part.element

abstract_ids = [
    int(item.get(qn("w:abstractNumId")))
    for item in numbering.findall(qn("w:abstractNum"))
]
abstract_id = max(abstract_ids or [0]) + 1

abstract_num = OxmlElement("w:abstractNum")
abstract_num.set(qn("w:abstractNumId"), str(abstract_id))

multi = OxmlElement("w:multiLevelType")
multi.set(qn("w:val"), "multilevel")
abstract_num.append(multi)

# Define the first numbering level
level = OxmlElement("w:lvl")
level.set(qn("w:ilvl"), "0")

start = OxmlElement("w:start")
start.set(qn("w:val"), "1")
level.append(start)

num_fmt = OxmlElement("w:numFmt")
num_fmt.set(qn("w:val"), "decimal")
level.append(num_fmt)

level_text = OxmlElement("w:lvlText")
level_text.set(qn("w:val"), "%1.")
level.append(level_text)

alignment = OxmlElement("w:lvlJc")
alignment.set(qn("w:val"), "left")
level.append(alignment)

# Set numbering indentation
ppr = OxmlElement("w:pPr")
ind = OxmlElement("w:ind")
ind.set(qn("w:left"), "720")
ind.set(qn("w:hanging"), "360")
ppr.append(ind)
level.append(ppr)

abstract_num.append(level)
numbering.append(abstract_num)

# Create a numbering instance
num_ids = [
    int(item.get(qn("w:numId")))
    for item in numbering.findall(qn("w:num"))
]
num_id = max(num_ids or [0]) + 1

num = OxmlElement("w:num")
num.set(qn("w:numId"), str(num_id))

abstract_ref = OxmlElement("w:abstractNumId")
abstract_ref.set(qn("w:val"), str(abstract_id))
num.append(abstract_ref)
numbering.append(num)

# Add the user's text as Heading 1
paragraph = doc.add_paragraph(style="Heading 1")
paragraph.add_run(text)

# Apply automatic numbering
paragraph_ppr = paragraph._p.get_or_add_pPr()
num_pr = OxmlElement("w:numPr")

ilvl = OxmlElement("w:ilvl")
ilvl.set(qn("w:val"), "0")

num_id_tag = OxmlElement("w:numId")
num_id_tag.set(qn("w:val"), str(num_id))

num_pr.append(ilvl)
num_pr.append(num_id_tag)
paragraph_ppr.append(num_pr)

# Save the document
filename = "Formatted_Heading.docx"
doc.save(filename)

print(f"Document saved as {filename}")
