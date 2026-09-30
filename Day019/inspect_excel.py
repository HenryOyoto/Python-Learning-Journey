import openpyxl 
from pathlib import Path

folder = Path("tender_source_files")

files = folder.glob("*.xlsx")

all_data = []

record_id = 1

for file in files:
    workbook = openpyxl.load_workbook(file)
    sheet = workbook["Tender Data"]


    for row in range(2, sheet.max_row - 1):
        row_data = []


        for column in range(1, sheet.max_column + 1):
            value = sheet.cell(row, column).value
            row_data.append(value)

        row_data.append(file.name)
        row_data.append(f"MASTER-{record_id:06d}")

        all_data.append(row_data)

        record_id += 1
        
print(len(all_data))

print(all_data[0])
print(all_data[-1])



